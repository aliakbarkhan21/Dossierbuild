"""Import-layer self-check. Run with:  python scripts/check_import.py

Covers the LinkedIn export reader, resume text extraction, and the merge
planner. The Gemini call itself is not exercised here -- it needs a key and a
network -- but the conversion from the model's answer into the real schema is,
since that is where a bad response would actually cause damage.
"""

from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from dossier.ai import parse as ai_parse
from dossier.ingest import extract, linkedin
from dossier.ingest.merge import apply_merge_plan, build_merge_plan
from dossier.core.schema import Experience, Profile, SkillGroup, TextBlock, format_date

from _harness import check, run


def make_export(**overrides: str) -> bytes:
    files = {
        "Profile.csv": (
            "First Name,Last Name,Headline,Summary,Geo Location\n"
            "Ali,Khan,CS and AI student,Builds small tools end to end.,Manchester UK\n"
        ),
        "Positions.csv": (
            "Company Name,Title,Description,Location,Started On,Finished On\n"
            'Northgate Labs,Software Engineering Intern,'
            '"Cut nightly ETL runtime from 42 to 9 minutes\n'
            'Added 61 pytest cases over the invoice parser",Leeds,Jun 2025,Sep 2025\n'
        ),
        "Education.csv": (
            "School Name,Start Date,End Date,Notes,Degree Name,Field Of Study\n"
            "University of Manchester,2024,2027,,BSc,Computer Science with AI\n"
        ),
        "Skills.csv": "Name\nPython\nSQL\nPyTorch\n",
        "Certifications.csv": (
            "Name,Url,Authority,Started On\n"
            "Deep Learning Specialization,https://x.test/1,DeepLearning.AI,Mar 2025\n"
        ),
        "Honors.csv": "Title,Description,Issued On\nHackathon runner-up,2nd of 60 teams,Feb 2025\n",
        "Connections.csv": "First Name\nIgnored\n",
    }
    files.update(overrides)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        for name, body in files.items():
            archive.writestr(name, body)
    return buf.getvalue()


# --------------------------------------------------------------------------
# LinkedIn export
# --------------------------------------------------------------------------


@check("LinkedIn dates parse at the precision actually given")
def _() -> None:
    cases = {
        "Jun 2025": "2025-06",
        "Sept 2024": "2024-09",
        "06/2025": "2025-06",
        "2025-06-01": "2025-06",
        "2025": "2025",  # year only stays a year -- no invented month
        "": None,
        "sometime last year": None,
    }
    for raw, expected in cases.items():
        got = linkedin.parse_date(raw)
        assert got == expected, f"{raw!r} gave {got!r}, expected {expected!r}"


@check("a year-only education date is never turned into January")
def _() -> None:
    profile = linkedin.parse_export(make_export()).profile
    entry = profile.education[0]
    assert entry.start == "2024", entry.start
    assert entry.end == "2027", entry.end
    assert format_date(entry.end) == "2027", "should render as a bare year"


@check("a full export parses every section it should")
def _() -> None:
    result = linkedin.parse_export(make_export())
    p = result.profile
    assert p.basics.name == "Ali Khan", p.basics.name
    assert len(p.experience) == 1
    assert len(p.experience[0].bullets) == 2, "description should split into 2 bullets"
    assert len(p.education) == 1
    assert len(p.certifications) == 1
    assert len(p.awards) == 1


@check("single-column Skills.csv is not swallowed")
def _() -> None:
    # This one is here because an earlier version of the CSV reader dropped
    # every leading line without a comma, which emptied this file entirely.
    result = linkedin.parse_export(make_export())
    assert result.profile.skills, "Skills.csv produced no skills"
    assert result.profile.skills[0].items == ["Python", "SQL", "PyTorch"]


@check("a Notes: preamble before the header row is skipped")
def _() -> None:
    result = linkedin.parse_export(make_export(**{"Skills.csv": "Notes:\nName\nRust\nGo\n"}))
    assert result.profile.skills[0].items == ["Rust", "Go"]


@check("a non-ZIP upload is refused with an explanation")
def _() -> None:
    try:
        linkedin.parse_export(b"this is not a zip file")
    except ValueError as exc:
        assert "ZIP" in str(exc)
    else:
        raise AssertionError("random bytes should not parse")


@check("a ZIP with no recognised CSVs is refused")
def _() -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        archive.writestr("Messages.csv", "From,To\na,b\n")
    try:
        linkedin.parse_export(buf.getvalue())
    except ValueError as exc:
        assert "Positions" in str(exc)
    else:
        raise AssertionError("an unrelated archive should be refused")


@check("bullet glyphs and numbering are stripped from descriptions")
def _() -> None:
    blocks = linkedin.split_description("- Built a thing\n2. Shipped it\n• Measured it")
    assert [b.text for b in blocks] == ["Built a thing", "Shipped it", "Measured it"]


# --------------------------------------------------------------------------
# Merge
# --------------------------------------------------------------------------


@check("an entry already in the profile is flagged as a duplicate, not re-added")
def _() -> None:
    current = Profile.empty()
    current.experience.append(
        Experience(role="Software Engineering Intern", organisation="Northgate Labs")
    )
    imported = linkedin.parse_export(make_export()).profile
    plan = build_merge_plan(current, imported, "test")

    dupes = [c for c in plan.candidates if c.section == "experience" and c.is_duplicate]
    assert dupes, "the matching internship should be flagged"
    apply_merge_plan(current, plan, set(), plan.default_selection())
    assert len(current.experience) == 1, "a duplicate should not be added by default"


@check("duplicate detection ignores dates and edited bullets")
def _() -> None:
    current = Profile.empty()
    current.experience.append(
        Experience(
            role="software engineering intern",  # different case
            organisation="Northgate  Labs",  # different spacing
            start="2020-01",  # different dates
            bullets=[TextBlock(text="Completely different wording")],
        )
    )
    imported = linkedin.parse_export(make_export()).profile
    plan = build_merge_plan(current, imported, "test")
    dupes = [c for c in plan.candidates if c.section == "experience" and c.is_duplicate]
    assert dupes, "identity should be role + organisation only"


@check("import never overwrites a filled field unless it is ticked")
def _() -> None:
    current = Profile.empty()
    current.basics.headline = "My own carefully written headline"
    imported = linkedin.parse_export(make_export()).profile
    plan = build_merge_plan(current, imported, "test")

    headline = next(f for f in plan.fields if f.path == "basics.headline")
    assert headline.conflicts, "a filled field should be marked as conflicting"

    apply_merge_plan(current, plan, set(), set())  # nothing ticked
    assert current.basics.headline == "My own carefully written headline"

    apply_merge_plan(current, plan, {"basics.headline"}, set())  # ticked
    assert current.basics.headline == "CS and AI student"


@check("skill groups merge by name, case-insensitively, without duplicating items")
def _() -> None:
    current = Profile.empty()
    current.skills.append(SkillGroup(label="Imported from LinkedIn", items=["python", "Rust"]))
    imported = linkedin.parse_export(make_export()).profile
    plan = build_merge_plan(current, imported, "test")
    apply_merge_plan(current, plan, set(), {c.key for c in plan.candidates})

    assert len(current.skills) == 1, "a same-named group should merge, not duplicate"
    items = current.skills[0].items
    assert items == ["python", "Rust", "SQL", "PyTorch"], items


@check("applying a plan produces ids that are still unique")
def _() -> None:
    from dossier.core.schema import all_ids

    current = Profile.empty()
    imported = linkedin.parse_export(make_export()).profile
    plan = build_merge_plan(current, imported, "test")
    apply_merge_plan(current, plan, {f.path for f in plan.fields}, plan.default_selection())
    ids = all_ids(current)
    assert len(ids) == len(set(ids)), "merge introduced duplicate ids"


# --------------------------------------------------------------------------
# Extraction
# --------------------------------------------------------------------------


@check("plain text extraction tidies ligatures and collapses blank lines")
def _() -> None:
    raw = "Proﬁle\n\n\n\nBuilt   a thing\n• bullet\n".encode("utf-8")
    result = extract.extract(raw, "cv.txt")
    assert "Profile" in result.text, result.text
    assert "\n\n\n" not in result.text
    assert "Built a thing" in result.text


@check("an unsupported file type is refused by name")
def _() -> None:
    for name, expect in [("resume.doc", ".docx"), ("resume.pages", "Supported")]:
        try:
            extract.extract(b"x", name)
        except ValueError as exc:
            assert expect in str(exc), f"{name}: {exc}"
        else:
            raise AssertionError(f"{name} should be refused")


@check("a DOCX round-trips, including text inside layout tables")
def _() -> None:
    try:
        import docx
    except ImportError:
        raise AssertionError("python-docx not installed; run pip install -r requirements.txt")

    document = docx.Document()
    document.add_paragraph("Ali Khan")
    table = document.add_table(rows=1, cols=1)
    table.rows[0].cells[0].paragraphs[0].text = "Cut ETL runtime from 42 to 9 minutes"
    buf = io.BytesIO()
    document.save(buf)

    result = extract.extract(buf.getvalue(), "cv.docx")
    assert "Ali Khan" in result.text
    assert "Cut ETL runtime" in result.text, "table cell text should be extracted"


# --------------------------------------------------------------------------
# AI conversion (no network)
# --------------------------------------------------------------------------


@check("the model's answer converts into a profile the real schema accepts")
def _() -> None:
    raw = ai_parse.RawResume(
        name="Ali Khan",
        experience=[
            ai_parse.RawExperience(
                role="Intern",
                organisation="Northgate",
                employment_type="internship",
                start="Jun 2025",
                end="",
                bullets=["Cut ETL runtime from 42 to 9 minutes", "   ", ""],
            )
        ],
        skills=[ai_parse.RawSkillGroup(label="", items=["Python"]),
                ai_parse.RawSkillGroup(label="Empty", items=[])],
    )
    profile = ai_parse.to_profile(raw)
    Profile.model_validate(profile.model_dump())  # must satisfy the real contract

    entry = profile.experience[0]
    assert entry.start == "2025-06", entry.start
    assert entry.end is None, "an empty end date means ongoing"
    assert len(entry.bullets) == 1, "blank bullets should be dropped"
    assert entry.employment_type == "Internship"
    assert [g.label for g in profile.skills] == ["Skills"], "unlabelled group gets a default name"


@check("the model returning a bare year is preserved, not padded to January")
def _() -> None:
    raw = ai_parse.RawResume(
        education=[ai_parse.RawEducation(institution="X", start="2024", end="2027")]
    )
    entry = ai_parse.to_profile(raw).education[0]
    assert (entry.start, entry.end) == ("2024", "2027"), (entry.start, entry.end)


@check("a date the model states in words is kept rather than dropped")
def _() -> None:
    # This used to return None for anything unparseable, so a CV reading
    # "Summer 2024" imported with no date at all -- worse than an odd-looking
    # one, because nothing on screen said a fact had been dropped.
    raw = ai_parse.RawResume(
        education=[ai_parse.RawEducation(institution="X", start="2024-09", end="whenever")]
    )
    profile = ai_parse.to_profile(raw)
    assert profile.education[0].start == "2024-09", "a real date is still normalised"
    assert profile.education[0].end == "whenever", "words are kept"


@check("a hyphen a PDF broke across lines is rejoined before the model sees it")
def _() -> None:
    from dossier.ingest.extract import _tidy

    assert _tidy("driving large-\nscale change") == "driving large-scale change"
    # A line break that is not inside a word is left alone: it is the main
    # signal separating one bullet from the next.
    assert _tidy("first line\nsecond line") == "first line\nsecond line"


@check("an unrecognised employment type falls back to Other instead of failing")
def _() -> None:
    raw = ai_parse.RawResume(
        experience=[ai_parse.RawExperience(role="Intern", employment_type="Summer Analyst")]
    )
    assert ai_parse.to_profile(raw).experience[0].employment_type == "Other"


@check("ids come from us, never from the model")
def _() -> None:
    fields = set(ai_parse.RawResume.model_fields)
    assert "id" not in fields
    for model in (ai_parse.RawExperience, ai_parse.RawProject, ai_parse.RawEducation):
        assert "id" not in model.model_fields, f"{model.__name__} should not expose an id"


def main() -> int:
    return run(__name__)


if __name__ == "__main__":
    raise SystemExit(main())
