"""Getting plain text out of an uploaded resume file.

This half of the resume import is deterministic -- no model involved. Keeping
it separate from ``ai_parse`` means the failure modes stay distinguishable: if
the text below looks wrong, the problem is extraction; if the text looks right
but the parsed profile is wrong, the problem is the model.

A caveat worth knowing about PDFs: a PDF does not contain paragraphs, columns
or reading order. It contains glyphs with coordinates. Extraction re-guesses
the reading order from those positions, which is why a two-column resume often
comes out interleaved. That is a property of the format, not a bug here -- and
it is exactly why the text is shown for review before anything is parsed.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass

SUPPORTED_SUFFIXES = (".pdf", ".docx", ".txt", ".md")


class ExtractError(ValueError):
    """This file cannot be read, with a sentence saying why.

    A subclass of ValueError so existing callers that catch ValueError keep
    working, and a named type so the HTTP layer can answer 422 with the
    message rather than 500 with a traceback.
    """


@dataclass
class Extraction:
    text: str
    source_name: str
    kind: str
    pages: int = 0
    warnings: list[str] = None  # type: ignore[assignment]

    def __post_init__(self) -> None:
        if self.warnings is None:
            self.warnings = []

    @property
    def char_count(self) -> int:
        return len(self.text)

    @property
    def looks_empty(self) -> bool:
        return len(self.text.strip()) < 120


def _tidy(text: str) -> str:
    """Normalise whitespace without destroying line structure.

    Line breaks matter -- they are the main signal that separates one bullet
    from the next -- so runs of blank lines collapse to one, but single breaks
    survive. Ligatures and the non-breaking hyphens PDF producers like are
    folded back to ASCII so downstream matching is not tripped by them.
    """
    replacements = {
        "ﬁ": "fi", "ﬂ": "fl", " ": " ", "‐": "-", "‑": "-",
        "–": "-", "—": "-", "‘": "'", "’": "'",
        "“": '"', "”": '"', "•": "-", "­": "",
    }
    for bad, good in replacements.items():
        text = text.replace(bad, good)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    # A hyphen at the end of a line is a compound word the layout happened to
    # wrap, and the newline after it is a fact about the page rather than
    # about the words. Left in, it reaches the model as "large-\nscale" and
    # comes back as "large- scale", which then prints on the new resume --
    # observed on a real import, three times in one CV.
    #
    # The hyphen is kept rather than dropped. Dropping it is right for a
    # typesetter's hyphenation ("initia-tives") and wrong for the compounds a
    # CV is full of, and Word does not hyphenate by default -- so on a resume
    # a hyphen before a line break is nearly always a hyphen the writer typed.
    text = re.sub(r"-\n(?=[A-Za-z])", "-", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_pdf(data: bytes, name: str) -> Extraction:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("pypdf is not installed. Run: pip install pypdf") from exc

    reader = PdfReader(io.BytesIO(data))
    warnings: list[str] = []

    if reader.is_encrypted:
        try:
            reader.decrypt("")  # many resumes are "encrypted" with an empty password
        except Exception as exc:  # noqa: BLE001
            raise ExtractError(
                "That PDF is password protected. Save an unprotected copy and upload that."
            ) from exc

    pages = [page.extract_text() or "" for page in reader.pages]
    text = _tidy("\n\n".join(pages))

    if len(text.strip()) < 120:
        warnings.append(
            "Almost no text came out. This is usually a scanned or image-based PDF -- "
            "the page is a picture of a resume, so there are no characters to read. "
            "Export a text PDF from the original document, or paste the text in directly."
        )

    return Extraction(text=text, source_name=name, kind="pdf", pages=len(reader.pages), warnings=warnings)


def extract_docx(data: bytes, name: str) -> Extraction:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("python-docx is not installed. Run: pip install python-docx") from exc

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]

    # Plenty of resumes lay themselves out in an invisible table. Those
    # paragraphs are not in document.paragraphs, so read the cells too.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)

    return Extraction(text=_tidy("\n".join(parts)), source_name=name, kind="docx")


def extract_text_file(data: bytes, name: str) -> Extraction:
    return Extraction(text=_tidy(data.decode("utf-8", errors="replace")), source_name=name, kind="text")


def extract(data: bytes, filename: str) -> Extraction:
    """Dispatch on file extension."""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return extract_pdf(data, filename)
    if lower.endswith(".docx"):
        return extract_docx(data, filename)
    if lower.endswith((".txt", ".md")):
        return extract_text_file(data, filename)
    if lower.endswith(".doc"):
        raise ExtractError(
            "Old-style .doc files are not readable here. Open it in Word or Google Docs "
            "and save as .docx or PDF."
        )
    raise ExtractError(f"Cannot read {filename}. Supported: {', '.join(SUPPORTED_SUFFIXES)}")
